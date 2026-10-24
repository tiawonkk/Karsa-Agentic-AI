import datetime
from pathlib import Path
import docx
from langchain_core.messages import HumanMessage

from backend.core.database import initialize_master_database
from backend.core.config import settings
from backend.modules.base_module import ModuleContext
from backend.modules.study_daily.tools import (
    parse_due_date_natural,
    is_auto_urgent,
    add_task_in_db,
    list_tasks_from_db,
    complete_task_in_db,
    drop_task_in_db,
    extract_text_from_file,
    generate_skill_roadmap,
)
from backend.modules.study_daily.module import StudyDailyModule


def test_parse_due_date_natural():
    d_besok = parse_due_date_natural("kumpul besok malam")
    assert d_besok is not None
    assert d_besok.hour == 21

    d_lusa = parse_due_date_natural("tenggat lusa")
    assert d_lusa is not None

    d_iso = parse_due_date_natural("deadline 2026-11-15 14:00")
    assert d_iso is not None
    assert d_iso.year == 2026
    assert d_iso.month == 11
    assert d_iso.day == 15
    assert d_iso.hour == 14


def test_is_auto_urgent():
    now = datetime.datetime.now()
    due_24h = now + datetime.timedelta(hours=24)
    assert is_auto_urgent(due_24h) == 1

    due_72h = now + datetime.timedelta(hours=72)
    assert is_auto_urgent(due_72h) == 0

    assert is_auto_urgent(None) == 0


def test_eisenhower_quadrants_and_tasks_db(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    (tmp_path / "storage").mkdir(parents=True, exist_ok=True)
    initialize_master_database(settings.master_db_path)

    now = datetime.datetime.now()

    # 1. Q1_DO: Important=1, Urgent=1 (deadline tomorrow)
    t1 = add_task_in_db(
        title="Tugas Besar Kalkulus",
        course="Kalkulus II",
        due_at=now + datetime.timedelta(hours=20),
        is_important=1,
    )
    assert t1["quadrant"] == "Q1_DO"
    assert t1["is_urgent"] == 1
    assert t1["is_important"] == 1

    # 2. Q2_SCHEDULE: Important=1, Urgent=0 (deadline 7 days)
    t2 = add_task_in_db(
        title="Proposal Skripsi",
        course="Metopen",
        due_at=now + datetime.timedelta(days=7),
        is_important=1,
    )
    assert t2["quadrant"] == "Q2_SCHEDULE"
    assert t2["is_urgent"] == 0

    # 3. Complete task
    done_t1 = complete_task_in_db(t1["id"])
    assert done_t1["status"] == "done"
    assert done_t1["completed_at"] is not None

    # 4. List tasks filter
    todo_list = list_tasks_from_db(status="todo")
    assert len(todo_list) == 1
    assert todo_list[0]["id"] == t2["id"]


def test_extract_text_from_file(tmp_path: Path):
    # Text file
    txt_file = tmp_path / "notes.txt"
    txt_file.write_text("Catatan kuliah Kecerdasan Buatan.", encoding="utf-8")
    assert "Kecerdasan Buatan" in extract_text_from_file(txt_file)

    # Word .docx file
    docx_file = tmp_path / "kuliah.docx"
    doc = docx.Document()
    doc.add_heading("Algoritma Pemrograman", level=1)
    doc.add_paragraph("Materi pertemuan 5 tentang Graph dan Tree.")
    doc.save(str(docx_file))

    extracted = extract_text_from_file(docx_file)
    assert "Algoritma Pemrograman" in extracted
    assert "Graph dan Tree" in extracted


def test_generate_skill_roadmap(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    (tmp_path / "storage").mkdir(parents=True, exist_ok=True)
    initialize_master_database(settings.master_db_path)

    res = generate_skill_roadmap("python data science", register_as_tasks=True)
    assert res["matched_curriculum"] == "python_data_science"
    assert len(res["roadmap"]["weeks"]) == 4
    assert len(res["tasks_registered"]) == 4

    tasks = list_tasks_from_db(status="todo")
    assert len(tasks) == 4


def test_study_daily_subgraph_flow(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    (tmp_path / "storage").mkdir(parents=True, exist_ok=True)
    initialize_master_database(settings.master_db_path)

    mod = StudyDailyModule()
    ctx = ModuleContext(session_thread_id="test_study_thread_01")
    graph = mod.build_graph(ctx)

    # 1. Fast-path add task
    res1 = graph.invoke({"messages": [HumanMessage(content="ingatkan tugas fisika dasar deadline besok malam")]})
    assert len(res1["messages"]) > 0
    content1 = str(res1["messages"][-1].content)
    assert "Tugas Berhasil Ditambahkan" in content1
    assert "Eisenhower Matrix" in content1

    # 2. List tasks
    res2 = graph.invoke({"messages": [HumanMessage(content="daftar tugas kuliah saya")]})
    content2 = str(res2["messages"][-1].content)
    assert "Eisenhower Matrix" in content2
    assert "fisika dasar" in content2.lower()

    # 3. Roadmap query
    res3 = graph.invoke({"messages": [HumanMessage(content="buatkan roadmap belajar web development")]})
    content3 = str(res3["messages"][-1].content)
    assert "Modern Web Development" in content3
    assert "Minggu 1" in content3

    # 4. Complete task
    res4 = graph.invoke({"messages": [HumanMessage(content="tugas 1 selesai")]})
    content4 = str(res4["messages"][-1].content)
    assert "berhasil ditandai selesai" in content4


def test_extract_text_from_pptx(tmp_path: Path):
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Hukum Newton"
    slide.placeholders[1].text = "F = m a"
    f = tmp_path / "materi.pptx"
    prs.save(str(f))

    text = extract_text_from_file(f)
    assert "[Slide 1]" in text
    assert "Hukum Newton" in text and "F = m a" in text
