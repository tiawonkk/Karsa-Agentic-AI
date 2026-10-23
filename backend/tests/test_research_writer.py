from pathlib import Path
import docx
import pytest
from langchain_core.messages import HumanMessage

from backend.core.database import initialize_master_database, get_db_cursor
from backend.core.config import settings
from backend.modules.base_module import ModuleContext
from backend.modules.research_writer.research_tools import (
    search_and_cache_papers,
    extract_paper_parameters,
)
from backend.modules.research_writer.writer_tools import (
    generate_imrad_template,
    compile_imrad_to_docx,
)
from backend.modules.research_writer.module import ResearchModule, WriterModule


def test_search_and_cache_papers(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    (tmp_path / "storage").mkdir(parents=True, exist_ok=True)
    initialize_master_database(settings.master_db_path)

    papers = search_and_cache_papers("agentic ai systems", source="arxiv", max_results=2)
    assert len(papers) > 0
    assert "title" in papers[0]
    assert "authors" in papers[0]

    with get_db_cursor(settings.master_db_path) as cursor:
        cursor.execute("SELECT count(*) FROM academic_papers;")
        count = cursor.fetchone()[0]
        assert count > 0


def test_extract_paper_parameters(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    (tmp_path / "storage").mkdir(parents=True, exist_ok=True)
    initialize_master_database(settings.master_db_path)

    res = extract_paper_parameters(
        paper_title="Agentic AI Student Companion",
        abstract_text="We introduce a deep learning transformer architecture with empirical evaluation on benchmark datasets.",
    )
    assert res["publication_year"] == 2025
    assert "dataset_source" in res
    assert "methodology" in res
    assert "limitations" in res
    assert "evaluation_metrics" in res


def test_generate_imrad_and_compile_docx(tmp_path: Path):
    imrad = generate_imrad_template(topic="Deteksi Plagiarisme Kode", language="id")
    assert "ABSTRAK" in imrad["abstract"] or len(imrad["abstract"]) > 50
    assert "METODOLOGI" in imrad["methods"] or len(imrad["methods"]) > 50
    assert len(imrad["references"]) >= 3

    out_file = tmp_path / "Test_Paper.docx"
    comp_res = compile_imrad_to_docx(imrad, out_file, overwrite=False)
    assert comp_res["status"] == "success"
    assert out_file.exists()
    assert out_file.stat().st_size > 0

    # Read back using python-docx to verify document structure
    doc = docx.Document(str(out_file))
    paragraphs_text = [p.text for p in doc.paragraphs if p.text]
    assert any("PENDAHULUAN" in p for p in paragraphs_text)
    assert any("METODOLOGI" in p for p in paragraphs_text)
    assert any("HASIL DAN ANALISIS" in p for p in paragraphs_text)
    assert any("DAFTAR PUSTAKA" in p for p in paragraphs_text)

    # Tables verification
    assert len(doc.tables) >= 1
    assert "Metrik Evaluasi" in doc.tables[0].rows[0].cells[0].text

    # Test overwrite protection raises FileExistsError
    with pytest.raises(FileExistsError):
        compile_imrad_to_docx(imrad, out_file, overwrite=False)


def test_research_subgraph_flow(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    (tmp_path / "storage").mkdir(parents=True, exist_ok=True)
    initialize_master_database(settings.master_db_path)

    mod = ResearchModule()
    ctx = ModuleContext(session_thread_id="test_res_thread")
    graph = mod.build_graph(ctx)

    # 1. Search paper
    res1 = graph.invoke({"messages": [HumanMessage(content="carikan paper arxiv tentang agentic ai")]})
    assert len(res1["messages"]) > 0
    content1 = str(res1["messages"][-1].content)
    assert "Hasil Pencarian Jurnal Ilmiah" in content1
    assert "arxiv" in content1.lower()

    # 2. Extract parameters
    res2 = graph.invoke({"messages": [HumanMessage(content="ekstrak 5 parameter paper transformer attention")]})
    content2 = str(res2["messages"][-1].content)
    assert "Ekstraksi 5 Parameter Kunci Riset" in content2
    assert "Metodologi Riset" in content2


def test_writer_subgraph_flow(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    (tmp_path / "storage").mkdir(parents=True, exist_ok=True)
    initialize_master_database(settings.master_db_path)

    mod = WriterModule()
    ctx = ModuleContext(session_thread_id="test_writer_thread")
    graph = mod.build_graph(ctx)

    # 1. Draft IMRaD
    res1 = graph.invoke({"messages": [HumanMessage(content="buatkan draft paper imrad tentang multi agent student companion")]})
    assert len(res1["messages"]) > 0
    content1 = str(res1["messages"][-1].content)
    assert "Draf Naskah Ilmiah IMRaD" in content1
    assert "PENDAHULUAN" in content1
    assert "METODOLOGI" in content1

    # 2. Compile to DOCX
    res2 = graph.invoke({"messages": [HumanMessage(content="kompilasi ke file docx paper sistem karsa")]})
    content2 = str(res2["messages"][-1].content)
    assert "Dokumen Microsoft Word (.docx) Berhasil Dikompilasi" in content2
    assert "Nama File" in content2
