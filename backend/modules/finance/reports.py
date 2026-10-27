"""Monthly finance report (categories, savings ratio, mutations) with CSV and .docx export. All numbers come from SQLite."""
from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any

from backend.core.config import settings
from backend.core.database import get_db_cursor
from backend.modules.finance.tools import VALID_CATEGORIES, format_rupiah, get_budget_snapshot, get_current_period

PERIOD_RE = re.compile(r"\d{4}-(0[1-9]|1[0-2])")


class ReportError(ValueError):
    pass


def _check_period(period: str | None) -> str:
    period = period or get_current_period()
    if not PERIOD_RE.fullmatch(period):
        raise ReportError("Periode harus berformat YYYY-MM.")
    return period


def monthly_report(period: str | None = None, db_path: Path | str | None = None) -> dict[str, Any]:
    """Everything one month's report needs: category table, savings ratio, daily spend, and the full mutation list."""
    period = _check_period(period)
    snap = get_budget_snapshot(period=period, db_path=db_path)
    with get_db_cursor(db_path) as cursor:
        cursor.execute(
            "SELECT id, txn_date, amount_idr, category, merchant, description, source FROM financial_transactions "
            "WHERE period = ? ORDER BY txn_date, id;",
            (period,),
        )
        expenses = [dict(r) for r in cursor.fetchall()]
        cursor.execute(
            "SELECT id, txn_date, amount_idr, label, note FROM financial_income WHERE period = ? ORDER BY txn_date, id;",
            (period,),
        )
        incomes = [dict(r) for r in cursor.fetchall()]

    total_spent = sum(e["amount_idr"] for e in expenses)
    total_income = sum(i["amount_idr"] for i in incomes)
    # money available this month = the planned capital plus manual income; saved = what was not spent
    available = snap["total_capital_idr"] + total_income
    saved = available - total_spent
    saving_rate = round(max(saved, 0) / available * 100, 1) if available > 0 else None

    categories = []
    for cat in VALID_CATEGORIES:
        c = snap["categories"][cat]
        categories.append(
            {
                "category": cat,
                "spent": c["spent"],
                "allocated": c["allocated"],
                "share_pct": round(c["spent"] / total_spent * 100, 1) if total_spent else 0.0,
                "over_budget": c["spent"] > c["allocated"] > 0,
                "remaining": c["remaining"],
            }
        )

    daily: dict[str, int] = {}
    for e in expenses:
        day = str(e["txn_date"])[:10]
        daily[day] = daily.get(day, 0) + e["amount_idr"]
    top = sorted(expenses, key=lambda e: e["amount_idr"], reverse=True)[:5]
    mutations = sorted(
        [{"date": str(e["txn_date"])[:10], "type": "keluar", "category": e["category"], "label": e["merchant"] or e["description"] or "-", "amount": -e["amount_idr"]} for e in expenses]
        + [{"date": str(i["txn_date"])[:10], "type": "masuk", "category": "pemasukan", "label": i["label"], "amount": i["amount_idr"]} for i in incomes],
        key=lambda m: m["date"],
    )
    biggest = max(categories, key=lambda c: c["spent"]) if total_spent else None
    notes = []
    if not expenses and not incomes:
        notes.append("Belum ada transaksi pada periode ini.")
    if biggest and biggest["spent"]:
        notes.append(f"Pengeluaran terbesar: {biggest['category']} ({biggest['share_pct']}%).")
    over = [c["category"] for c in categories if c["over_budget"]]
    if over:
        notes.append("Melebihi alokasi: " + ", ".join(over) + ".")
    return {
        "period": period,
        "total_capital_idr": snap["total_capital_idr"],
        "total_income_idr": total_income,
        "total_spent_idr": total_spent,
        "saved_idr": saved,
        "saving_rate_pct": saving_rate,
        "daily_safe_spend_idr": snap["daily_safe_spend_idr"],
        "remaining_days": snap["remaining_days"],
        "categories": categories,
        "daily_spend": [{"date": d, "spent": v} for d, v in sorted(daily.items())],
        "top_expenses": top,
        "mutations": mutations,
        "notes": notes,
    }


def _out_path(period: str, ext: str) -> Path:
    settings.outputs_dir.mkdir(parents=True, exist_ok=True)
    return settings.outputs_dir / f"laporan_keuangan_{period}.{ext}"


def export_csv(period: str | None = None, db_path: Path | str | None = None) -> dict[str, Any]:
    rep = monthly_report(period, db_path)
    path = _out_path(rep["period"], "csv")
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["tanggal", "jenis", "kategori", "keterangan", "jumlah_idr"])
        for m in rep["mutations"]:
            w.writerow([m["date"], m["type"], m["category"], m["label"], m["amount"]])
    return {"path": str(path), "name": path.name, "rows": len(rep["mutations"])}


def export_docx(period: str | None = None, db_path: Path | str | None = None) -> dict[str, Any]:
    import docx

    rep = monthly_report(period, db_path)
    d = docx.Document()
    d.add_heading(f"Laporan Keuangan {rep['period']}", level=1)
    d.add_paragraph(
        f"Modal: {format_rupiah(rep['total_capital_idr'])} | Pemasukan: {format_rupiah(rep['total_income_idr'])} | "
        f"Pengeluaran: {format_rupiah(rep['total_spent_idr'])} | Sisa: {format_rupiah(rep['saved_idr'])}"
    )
    rate = rep["saving_rate_pct"]
    d.add_paragraph(f"Rasio tabungan: {rate}%" if rate is not None else "Rasio tabungan: belum bisa dihitung (belum ada modal atau pemasukan).")
    d.add_paragraph(f"Daily Safe Spend: {format_rupiah(rep['daily_safe_spend_idr'])} per hari ({rep['remaining_days']} hari tersisa).")
    for n in rep["notes"]:
        d.add_paragraph(n, style="List Bullet")
    d.add_heading("Per kategori", level=2)
    t = d.add_table(rows=1, cols=4)
    for i, h in enumerate(["Kategori", "Terpakai", "Alokasi", "Porsi"]):
        t.rows[0].cells[i].text = h
    for c in rep["categories"]:
        row = t.add_row().cells
        row[0].text = c["category"]
        row[1].text = format_rupiah(c["spent"])
        row[2].text = format_rupiah(c["allocated"])
        row[3].text = f"{c['share_pct']}%"
    d.add_heading("Mutasi", level=2)
    m = d.add_table(rows=1, cols=4)
    for i, h in enumerate(["Tanggal", "Jenis", "Keterangan", "Jumlah"]):
        m.rows[0].cells[i].text = h
    for x in rep["mutations"]:
        row = m.add_row().cells
        row[0].text, row[1].text, row[2].text = x["date"], x["type"], f"{x['category']}: {x['label']}"
        row[3].text = ("-" if x["amount"] < 0 else "+") + format_rupiah(abs(x["amount"]))
    path = _out_path(rep["period"], "docx")
    d.save(path)
    return {"path": str(path), "name": path.name, "rows": len(rep["mutations"])}
