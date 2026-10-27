"""One-click clean + export for model training and papers: ML-ready table, descriptive statistics and SVG histograms.

Deterministic pandas/numpy only (no LLM, no network). Every transformation is listed in the report so the user can
reproduce it: nothing is hidden. The original file is never modified; outputs are new files in the outputs folder.
"""
from __future__ import annotations

import html
import time
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

from backend.core.config import settings
from backend.modules.notebook.dataset_cleaner import CleanOptions, clean_dataframe, load_dataset
from backend.modules.notebook.store import NotebookError

ONE_HOT_MAX_UNIQUE = 15
MAX_CHART_COLUMNS = 6


def _unique_path(stem: str, suffix: str) -> Path:
    settings.outputs_dir.mkdir(parents=True, exist_ok=True)
    p = settings.outputs_dir / f"{stem}{suffix}"
    return p if not p.exists() else settings.outputs_dir / f"{stem}_{time.strftime('%Y%m%d_%H%M%S')}{suffix}"


def describe(df: pd.DataFrame) -> pd.DataFrame:
    """Descriptive statistics table (one row per column) for numeric and text columns."""
    rows = []
    for c in df.columns:
        s = df[c]
        base = {"kolom": str(c), "tipe": "angka" if pd.api.types.is_numeric_dtype(s) else "teks", "n_terisi": int(s.notna().sum()), "n_kosong": int(s.isna().sum())}
        if base["tipe"] == "angka" and s.notna().any():
            n = pd.to_numeric(s, errors="coerce").astype(float)
            base.update({"rata_rata": round(n.mean(), 4), "std": round(n.std(ddof=1), 4) if n.notna().sum() > 1 else None, "min": n.min(), "q1": n.quantile(0.25), "median": n.median(), "q3": n.quantile(0.75), "maks": n.max()})
        else:
            top = s.mode(dropna=True)
            base.update({"nilai_unik": int(s.nunique(dropna=True)), "paling_sering": str(top.iloc[0]) if not top.empty else None, "frekuensi_teratas": int((s == top.iloc[0]).sum()) if not top.empty else 0})
        rows.append(base)
    return pd.DataFrame(rows)


def histogram_svg(values: pd.Series, title: str, bins: int = 10) -> str:
    v = pd.to_numeric(values, errors="coerce").dropna().astype(float)
    w, h, pad = 360, 200, 28
    if v.empty:
        return f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}"><text x="10" y="20" font-size="12">{html.escape(title)}: tidak ada data</text></svg>'
    counts, edges = np.histogram(v, bins=min(bins, max(1, v.nunique())))
    peak = max(int(counts.max()), 1)
    bw = (w - 2 * pad) / len(counts)
    bars = []
    for i, c in enumerate(counts):
        bh = (h - 2 * pad) * int(c) / peak
        bars.append(f'<rect x="{pad + i * bw + 1:.1f}" y="{h - pad - bh:.1f}" width="{bw - 2:.1f}" height="{bh:.1f}" fill="#6366f1"><title>{edges[i]:g} - {edges[i + 1]:g}: {int(c)}</title></rect>')
    return (
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" font-family="sans-serif" font-size="10">'
        f'<text x="{w / 2}" y="14" text-anchor="middle" font-size="12" font-weight="bold">{html.escape(title)}</text>'
        + "".join(bars)
        + f'<line x1="{pad}" y1="{h - pad}" x2="{w - pad}" y2="{h - pad}" stroke="#444"/>'
        f'<text x="{pad}" y="{h - 12}">{edges[0]:g}</text><text x="{w - pad}" y="{h - 12}" text-anchor="end">{edges[-1]:g}</text>'
        f'<text x="{pad - 4}" y="{pad}" text-anchor="end">{peak}</text></svg>'
    )


def build_ml_ready(path: Path | str, target: str | None = None, scale: str = "standard", encode: str = "auto") -> dict[str, Any]:
    """Clean, impute (median / mode), encode categoricals and scale numerics. Returns the report with the written files.

    scale: none | standard (z-score) | minmax. encode: auto (one-hot up to 15 values, else label) | label.
    The target column (if given) is never scaled; a text target is label-encoded and its mapping is reported.
    """
    if scale not in ("none", "standard", "minmax"):
        raise NotebookError("scale harus none, standard, atau minmax.")
    if encode not in ("auto", "label"):
        raise NotebookError("encode harus auto atau label.")
    p = Path(path)
    cleaned, report = clean_dataframe(load_dataset(p), CleanOptions(fill_missing="none"))
    if cleaned.empty or cleaned.shape[1] == 0:
        raise NotebookError("Dataset kosong setelah dibersihkan.")
    target_col = None
    if target:
        match = [c for c in cleaned.columns if str(c) == target or str(c) == target.strip().lower().replace(" ", "_")]
        if not match:
            raise NotebookError(f"Kolom target '{target}' tidak ada. Kolom tersedia: {', '.join(map(str, cleaned.columns[:20]))}")
        target_col = match[0]

    stats = describe(cleaned)  # statistics describe the cleaned data BEFORE imputation/encoding/scaling
    steps: list[str] = []
    df = cleaned.copy()

    for c in df.columns:
        miss = int(df[c].isna().sum())
        if not miss:
            continue
        if pd.api.types.is_numeric_dtype(df[c]):
            df[c] = df[c].fillna(df[c].median())
            steps.append(f"{c}: {miss} nilai kosong diisi median")
        else:
            mode = df[c].mode(dropna=True)
            df[c] = df[c].fillna(mode.iloc[0] if not mode.empty else "tidak_diketahui")
            steps.append(f"{c}: {miss} nilai kosong diisi modus")

    mappings: dict[str, dict[str, int]] = {}
    text_cols = [c for c in df.columns if not pd.api.types.is_numeric_dtype(df[c])]
    one_hot = []
    for c in text_cols:
        vals = sorted(df[c].astype(str).unique())
        if c == target_col or encode == "label" or len(vals) > ONE_HOT_MAX_UNIQUE:
            mapping = {v: i for i, v in enumerate(vals)}
            df[c] = df[c].astype(str).map(mapping)
            mappings[str(c)] = mapping
            steps.append(f"{c}: label encoding ({len(vals)} nilai)")
        else:
            one_hot.append(c)
    if one_hot:
        df = pd.get_dummies(df, columns=one_hot, dtype=int)
        steps.append("One-hot encoding: " + ", ".join(map(str, one_hot)))

    scaled = []
    if scale != "none":
        for c in df.columns:
            if c == target_col or not pd.api.types.is_numeric_dtype(df[c]) or set(df[c].dropna().unique()) <= {0, 1}:
                continue
            col = df[c].astype(float)
            if scale == "standard":
                sd = col.std(ddof=0)
                df[c] = (col - col.mean()) / sd if sd else 0.0
            else:
                rng = col.max() - col.min()
                df[c] = (col - col.min()) / rng if rng else 0.0
            scaled.append(str(c))
        if scaled:
            steps.append(f"Skala {scale}: " + ", ".join(scaled))

    stem = p.stem
    ml_path = _unique_path(f"{stem}_ml", ".csv")
    df.to_csv(ml_path, index=False, encoding="utf-8")
    stats_path = _unique_path(f"{stem}_statistik", ".csv")
    stats.to_csv(stats_path, index=False, encoding="utf-8-sig")
    docx_path = _unique_path(f"{stem}_statistik", ".docx")
    _stats_docx(stats, stem, docx_path)

    charts = []
    for c in [c for c in cleaned.columns if pd.api.types.is_numeric_dtype(cleaned[c])][:MAX_CHART_COLUMNS]:
        cp = _unique_path(f"{stem}_hist_{c}", ".svg")
        cp.write_text(histogram_svg(cleaned[c], f"Distribusi {c}"), encoding="utf-8")
        charts.append(cp.name)

    return {
        "rows": int(len(df)),
        "columns": int(df.shape[1]),
        "clean_report": report,
        "steps": steps,
        "label_mappings": mappings,
        "target": str(target_col) if target_col else None,
        "files": {"ml_csv": ml_path.name, "stats_csv": stats_path.name, "stats_docx": docx_path.name, "charts": charts},
    }


def _stats_docx(stats: pd.DataFrame, name: str, path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading(f"Statistik deskriptif: {name}", level=1)
    cols = ["kolom", "tipe", "n_terisi", "n_kosong", "rata_rata", "std", "min", "median", "maks"]
    t = d.add_table(rows=1, cols=len(cols))
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for _, r in stats.iterrows():
        cells = t.add_row().cells
        for i, c in enumerate(cols):
            v = r.get(c)
            cells[i].text = "" if v is None or (isinstance(v, float) and np.isnan(v)) else (f"{v:.4g}" if isinstance(v, (float, np.floating)) else str(v))
    d.save(path)
